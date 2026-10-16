import openpyxl
import os
import glob

base_dir = r"D:\GRAVITY\Rabies\Rabies 2026  ข้อมูล ณ วันที่ 15 ตุลาคม 2569"

xlsx_files = glob.glob(os.path.join(base_dir, "**", "*.xlsx"), recursive=True)
print(f"Total xlsx files: {len(xlsx_files)}")

# Let's inspect each xlsx file to see how many rows of patients
total_extracted_rows = 0
for f in xlsx_files:
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        sheet = wb.active
        # Count rows with data
        row_count = 0
        for r in range(1, sheet.max_row + 1):
            val = sheet.cell(r, 2).value or sheet.cell(r, 1).value
            if val and not any(header in str(val) for header in ['ลำดับ', 'ชื่อ', 'ตาราง', 'ผู้สัมผัส']):
                row_count += 1
        rel_path = os.path.relpath(f, base_dir)
        print(f"{rel_path}: ~{row_count} rows (max_row={sheet.max_row})")
        total_extracted_rows += row_count
    except Exception as e:
        print(f"Error reading {f}: {e}")

print(f"Total extracted rows across all files: {total_extracted_rows}")
